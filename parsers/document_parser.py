from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

from fastapi import UploadFile


@dataclass(frozen=True)
class ParsedDocument:
    filename: str
    text: str
    content_type: str | None = None


async def parse_upload(file: UploadFile) -> ParsedDocument:
    content = await file.read()
    filename = file.filename or "uploaded-resume.txt"
    suffix = filename.lower().rsplit(".", 1)[-1]
    if suffix == "pdf":
        text = parse_pdf(content)
    elif suffix == "docx":
        text = parse_docx(content)
    else:
        text = content.decode("utf-8", errors="ignore")
    return ParsedDocument(filename=filename, text=clean_text(text), content_type=file.content_type)


def parse_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception:
        return content.decode("utf-8", errors="ignore")


def parse_docx(content: bytes) -> str:
    try:
        from docx import Document

        doc = Document(BytesIO(content))
        return "\n".join(paragraph.text for paragraph in doc.paragraphs)
    except Exception:
        return content.decode("utf-8", errors="ignore")


def clean_text(text: str) -> str:
    return "\n".join(line.strip() for line in text.replace("\x00", "").splitlines() if line.strip())

